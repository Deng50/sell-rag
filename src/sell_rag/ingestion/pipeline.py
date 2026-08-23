from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable, Iterable

from sell_rag.domain import Chunk, DocumentElement, Product, Role
from sell_rag.settings import Settings
from sell_rag.storage import Database


SUPPORTED = {".pdf", ".docx", ".xlsx", ".csv", ".json", ".txt", ".md", ".png", ".jpg", ".jpeg"}
INJECTION_PATTERNS = [
    re.compile(pattern, re.I)
    for pattern in [
        r"ignore\s+(all\s+)?previous\s+instructions",
        r"忽略(以上|之前|所有).{0,12}(指令|要求)",
        r"system\s*prompt",
        r"你现在必须.{0,30}(输出|执行)",
    ]
]


@dataclass
class ParsedElement:
    content_type: str
    text: str = ""
    title: str = ""
    section_path: list[str] = field(default_factory=list)
    page_no: int | None = None
    sheet_name: str | None = None
    cell_range: str | None = None
    bbox: list[float] | None = None
    table_json: list[list[Any]] | None = None
    image_ref: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


class FileValidator:
    def __init__(self, max_file_mb: int = 50):
        self.max_bytes = max_file_mb * 1024 * 1024

    def validate(self, path: Path) -> None:
        if not path.is_file():
            raise ValueError("上传目标不是文件")
        if path.suffix.lower() not in SUPPORTED:
            raise ValueError(f"不支持的文件类型: {path.suffix}")
        size = path.stat().st_size
        if size == 0:
            raise ValueError("不允许上传空文件")
        if size > self.max_bytes:
            raise ValueError(f"文件超过 {self.max_bytes // 1024 // 1024} MB 限制")
        with path.open("rb") as handle:
            head = handle.read(12)
        suffix = path.suffix.lower()
        if suffix == ".pdf" and not head.startswith(b"%PDF"):
            raise ValueError("PDF 文件签名不正确")
        if suffix in {".docx", ".xlsx"}:
            if not head.startswith(b"PK") or not zipfile.is_zipfile(path):
                raise ValueError("Office 文件签名不正确")
            self._check_zip(path)
        if suffix == ".png" and not head.startswith(b"\x89PNG"):
            raise ValueError("PNG 文件签名不正确")
        if suffix in {".jpg", ".jpeg"} and not head.startswith(b"\xff\xd8\xff"):
            raise ValueError("JPEG 文件签名不正确")

    @staticmethod
    def _check_zip(path: Path) -> None:
        with zipfile.ZipFile(path) as archive:
            total = sum(info.file_size for info in archive.infolist())
            compressed = max(1, sum(info.compress_size for info in archive.infolist()))
            if total > 500 * 1024 * 1024 or total / compressed > 200:
                raise ValueError("Office 文件解压规模异常")
            if any(".." in Path(info.filename).parts for info in archive.infolist()):
                raise ValueError("Office 文件包含不安全路径")

    @staticmethod
    def sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for block in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest()


def parse_header_footer_candidates(elements: list[ParsedElement]) -> list[ParsedElement]:
    """Remove repeated margin text while preserving page metadata."""
    pages = {item.page_no for item in elements if item.page_no is not None}
    if len(pages) < 3:
        return elements
    candidates: dict[str, set[int]] = {}
    for item in elements:
        if item.page_no is None or not item.bbox or not item.text.strip():
            continue
        top, bottom = item.bbox[1], item.bbox[3]
        if top <= 0.08 or bottom >= 0.92:
            normalized = re.sub(r"\d+", "#", re.sub(r"\s+", " ", item.text.strip().lower()))
            candidates.setdefault(normalized, set()).add(item.page_no)
    repeated = {text for text, page_set in candidates.items() if len(page_set) / len(pages) >= 0.60}
    result = []
    for item in elements:
        normalized = re.sub(r"\d+", "#", re.sub(r"\s+", " ", item.text.strip().lower()))
        if normalized in repeated and item.bbox and (item.bbox[1] <= 0.08 or item.bbox[3] >= 0.92):
            continue
        result.append(item)
    return result


class DocumentParser:
    """Docling-first parser with explicit, testable fallbacks."""

    def __init__(self, image_describer: Callable[[Path], str] | None = None):
        self.image_describer = image_describer

    def parse(self, path: Path) -> tuple[list[ParsedElement], list[Product]]:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            return self._xlsx(path)
        if suffix == ".csv":
            return self._csv(path)
        if suffix == ".json":
            return self._json(path)
        if suffix in {".txt", ".md"}:
            return self._plain(path), []
        if suffix in {".png", ".jpg", ".jpeg"}:
            description = self.image_describer(path) if self.image_describer else ""
            return [ParsedElement("image", text=description, image_ref=str(path), metadata={"ocr": not bool(description)})], []
        try:
            return self._docling(path), []
        except (ImportError, ModuleNotFoundError):
            if suffix == ".docx":
                return self._docx_fallback(path), []
            if suffix == ".pdf":
                return self._pdf_fallback(path), []
            raise

    def _docling(self, path: Path) -> list[ParsedElement]:
        from docling.document_converter import DocumentConverter

        document = DocumentConverter().convert(str(path)).document
        exported = document.export_to_dict()
        elements: list[ParsedElement] = []
        for order, item in enumerate(exported.get("texts", [])):
            provenance = (item.get("prov") or [{}])[0]
            bbox_raw = provenance.get("bbox") or {}
            bbox = None
            if bbox_raw:
                bbox = [float(bbox_raw.get(key, 0)) for key in ("l", "t", "r", "b")]
            label = str(item.get("label", "text"))
            content_type = "title" if "title" in label or "heading" in label else "text"
            elements.append(ParsedElement(
                content_type=content_type,
                text=str(item.get("text", "")).strip(),
                title=str(item.get("text", "")).strip() if content_type == "title" else "",
                page_no=provenance.get("page_no"), bbox=bbox, metadata={"docling_label": label, "order": order},
            ))
        for table in exported.get("tables", []):
            provenance = (table.get("prov") or [{}])[0]
            grid = table.get("data", {}).get("grid") or table.get("table_cells") or []
            markdown = self._table_markdown(grid)
            elements.append(ParsedElement(
                content_type="table", text=markdown, table_json=grid,
                page_no=provenance.get("page_no"), metadata={"docling": True},
            ))
        return parse_header_footer_candidates([item for item in elements if item.text or item.image_ref])

    def _docx_fallback(self, path: Path) -> list[ParsedElement]:
        from docx import Document

        doc = Document(str(path))
        result, section = [], []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            if style.lower().startswith("heading"):
                level_match = re.search(r"(\d+)", style)
                level = int(level_match.group(1)) if level_match else 1
                section = section[: level - 1] + [text]
                result.append(ParsedElement("title", text=text, title=text, section_path=list(section)))
            else:
                result.append(ParsedElement("text", text=text, section_path=list(section)))
        for table in doc.tables:
            grid = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            result.append(ParsedElement("table", text=self._table_markdown(grid), table_json=grid,
                                        section_path=list(section)))
        return result

    def _pdf_fallback(self, path: Path) -> list[ParsedElement]:
        from pypdf import PdfReader

        result = []
        for page_no, page in enumerate(PdfReader(str(path)).pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                result.append(ParsedElement("text", text=text, page_no=page_no))
        if not result:
            raise RuntimeError("PDF 没有可提取文本，请安装 Docling OCR 依赖")
        return result

    @staticmethod
    def _plain(path: Path) -> list[ParsedElement]:
        result, section = [], []
        for block in re.split(r"\n\s*\n", path.read_text(encoding="utf-8-sig", errors="replace")):
            text = block.strip()
            if not text:
                continue
            if re.match(r"^#{1,6}\s+", text):
                title = re.sub(r"^#{1,6}\s+", "", text).strip()
                section = [title]
                result.append(ParsedElement("title", text=title, title=title, section_path=list(section)))
            else:
                result.append(ParsedElement("text", text=text, section_path=list(section)))
        return result

    def _xlsx(self, path: Path) -> tuple[list[ParsedElement], list[Product]]:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
        elements, products = [], []
        for sheet_values in values.worksheets:
            sheet_formulas = formulas[sheet_values.title]
            rows = list(sheet_values.iter_rows(values_only=True))
            if not rows:
                continue
            header_idx = next((i for i, row in enumerate(rows[:10]) if sum(x not in (None, "") for x in row) >= 2), 0)
            canonical = {"sku", "name", "category", "price", "stock", "target_group", "description", "active"}
            second_headers = self._headers(rows[header_idx + 1]) if header_idx + 1 < len(rows) else []
            header_count = 2 if len(canonical & set(second_headers)) >= 2 else 1
            if header_count == 1:
                headers = self._headers(rows[header_idx])
            else:
                width = max(len(rows[header_idx]), len(rows[header_idx + 1]))
                headers = []
                for column in range(width):
                    values_for_header = [str(rows[row_index][column]).strip()
                                         for row_index in range(header_idx, header_idx + header_count)
                                         if column < len(rows[row_index]) and rows[row_index][column] not in (None, "")]
                    leaf = self._normalize_header(values_for_header[-1]) if values_for_header else ""
                    headers.append(leaf if leaf in canonical else "/".join(values_for_header))
            table = [list(row) for row in rows[header_idx:] if any(x not in (None, "") for x in row)]
            end_col = get_column_letter(max(len(row) for row in table))
            cell_range = f"A{header_idx + 1}:{end_col}{header_idx + len(table)}"
            elements.append(ParsedElement(
                "table", text=self._table_markdown(table), table_json=table,
                sheet_name=sheet_values.title, cell_range=cell_range,
                metadata={"merged_ranges": [str(x) for x in sheet_values.merged_cells.ranges]},
            ))
            data_start = header_idx + header_count
            for offset, row in enumerate(rows[data_start:], data_start + 1):
                record = {headers[i]: value for i, value in enumerate(row[: len(headers)]) if headers[i] and value not in (None, "")}
                if not record:
                    continue
                formulas_for_row = {
                    headers[i]: sheet_formulas.cell(offset, i + 1).value
                    for i in range(min(len(headers), len(row)))
                    if isinstance(sheet_formulas.cell(offset, i + 1).value, str)
                    and sheet_formulas.cell(offset, i + 1).value.startswith("=")
                }
                product = self._product(record, path, sheet_values.title, offset, formulas_for_row)
                if product:
                    products.append(product)
                    elements.append(ParsedElement(
                        "product", text=self._product_text(product), title=product.name,
                        sheet_name=sheet_values.title, cell_range=f"A{offset}:{end_col}{offset}",
                        metadata={"sku": product.sku},
                    ))
        return elements, products

    def _csv(self, path: Path) -> tuple[list[ParsedElement], list[Product]]:
        products = []
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        for index, row in enumerate(rows, 2):
            product = self._product({self._normalize_header(k): v for k, v in row.items()}, path, None, index, {})
            if product:
                products.append(product)
        return [ParsedElement("product", text=self._product_text(item), title=item.name,
                              cell_range=f"row:{i + 2}", metadata={"sku": item.sku})
                for i, item in enumerate(products)], products

    def _json(self, path: Path) -> tuple[list[ParsedElement], list[Product]]:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        rows = payload if isinstance(payload, list) else payload.get("products", [])
        products = [item for i, row in enumerate(rows, 1)
                    if (item := self._product({self._normalize_header(k): v for k, v in row.items()},
                                              path, None, i, {}))]
        return [ParsedElement("product", text=self._product_text(item), title=item.name,
                              metadata={"sku": item.sku}) for item in products], products

    @staticmethod
    def _headers(row: Iterable[Any]) -> list[str]:
        return [DocumentParser._normalize_header(str(value or "").strip()) for value in row]

    @staticmethod
    def _normalize_header(value: str) -> str:
        mapping = {
            "商品编号": "sku", "编号": "sku", "SKU": "sku", "商品名称": "name", "名称": "name",
            "类别": "category", "分类": "category", "单价": "price", "价格": "price",
            "库存": "stock", "数量": "stock", "适合人群": "target_group", "目标人群": "target_group",
            "特点": "description", "卖点": "description", "描述": "description", "状态": "active",
        }
        return mapping.get(value, value)

    def _product(self, record: dict[str, Any], path: Path, sheet: str | None, row: int,
                 formulas: dict[str, Any]) -> Product | None:
        name = str(record.get("name") or "").strip()
        if not name:
            return None
        sku = str(record.get("sku") or hashlib.sha256(f"{path.name}:{sheet}:{name}".encode()).hexdigest()[:12]).strip()
        price_text = re.sub(r"[^0-9.-]", "", str(record.get("price") or "0")) or "0"
        try:
            price = Decimal(price_text)
        except InvalidOperation:
            price = Decimal("0")
        stock_text = re.search(r"\d+", str(record.get("stock") or "0"))
        target = [x.strip() for x in re.split(r"[,，、;/]", str(record.get("target_group") or "")) if x.strip()]
        active_value = str(record.get("active", "true")).lower()
        return Product(
            sku=sku, name=name, category=str(record.get("category") or "未分类"), price=price,
            stock=int(stock_text.group()) if stock_text else 0, target_group=target,
            description=str(record.get("description") or ""), active=active_value not in {"false", "0", "下架"},
            attributes={"source": path.name, "sheet_name": sheet, "row_id": row, "formulas": formulas},
        )

    @staticmethod
    def _product_text(product: Product) -> str:
        return "\n".join(filter(None, [
            f"商品编号：{product.sku}", f"商品名称：{product.name}", f"类别：{product.category}",
            f"价格：{product.price}元", f"库存：{product.stock}",
            f"适合人群：{'、'.join(product.target_group)}" if product.target_group else "",
            f"特点：{product.description}" if product.description else "",
        ]))

    @staticmethod
    def _table_markdown(grid: list[list[Any]]) -> str:
        if not grid:
            return ""
        width = max(len(row) for row in grid)
        rows = [[str(cell or "").replace("|", "\\|") for cell in row] + [""] * (width - len(row)) for row in grid]
        return "\n".join([
            "| " + " | ".join(rows[0]) + " |",
            "| " + " | ".join(["---"] * width) + " |",
            *("| " + " | ".join(row) + " |" for row in rows[1:]),
        ])


class ParentChildChunker:
    def __init__(self, child_tokens: int = 300, parent_tokens: int = 900, overlap_ratio: float = 0.12):
        self.child_tokens = child_tokens
        self.parent_tokens = parent_tokens
        self.overlap_ratio = overlap_ratio

    def chunk(self, document_id: str, version: int, acl: list[Role], elements: list[DocumentElement]) -> list[Chunk]:
        chunks, chunk_index = [], 0
        for element in elements:
            text = element.text.strip()
            if not text:
                continue
            parent_pieces = self._windows(text, self.parent_tokens, 0)
            for parent_no, parent_text in enumerate(parent_pieces):
                parent_id = self._id(document_id, version, element.element_id, "p", parent_no, parent_text)
                common = dict(
                    document_id=document_id, document_version=version, title=element.title,
                    section_path=element.section_path, page_no=element.page_no,
                    sheet_name=element.sheet_name, cell_range=element.cell_range,
                    content_type=element.content_type, acl=acl,
                    prompt_injection=any(p.search(parent_text) for p in INJECTION_PATTERNS),
                    metadata={"element_id": element.element_id, **element.metadata},
                )
                chunks.append(Chunk(
                    chunk_id=parent_id, parent_id=None, level="parent", index=chunk_index,
                    text=parent_text, content_hash=hashlib.sha256(parent_text.encode()).hexdigest(), **common,
                ))
                chunk_index += 1
                overlap = max(1, int(self.child_tokens * self.overlap_ratio))
                for child_no, child_text in enumerate(self._windows(parent_text, self.child_tokens, overlap)):
                    child_id = self._id(document_id, version, parent_id, "c", child_no, child_text)
                    chunks.append(Chunk(
                        chunk_id=child_id, parent_id=parent_id, level="child", index=chunk_index,
                        text=child_text, content_hash=hashlib.sha256(child_text.encode()).hexdigest(), **common,
                    ))
                    chunk_index += 1
        return chunks

    @staticmethod
    def _tokens(text: str) -> list[str]:
        return re.findall(r"[\u4e00-\u9fff]|[A-Za-z0-9_.+-]+|[^\s]", text)

    def _windows(self, text: str, size: int, overlap: int) -> list[str]:
        tokens = self._tokens(text)
        if len(tokens) <= size:
            return [text]
        result, start = [], 0
        while start < len(tokens):
            end = min(start + size, len(tokens))
            result.append("".join(tokens[start:end]))
            if end == len(tokens):
                break
            start = max(start + 1, end - overlap)
        return result

    @staticmethod
    def _id(document_id: str, version: int, scope: str, level: str, index: int, text: str) -> str:
        value = f"{document_id}:{version}:{scope}:{level}:{index}:{text}"
        return hashlib.sha256(value.encode()).hexdigest()[:32]


class IngestionService:
    PARSER_VERSION = "sell-rag-v1"

    def __init__(self, settings: Settings, database: Database,
                 image_describer: Callable[[Path], str] | None = None):
        self.settings = settings
        self.database = database
        self.validator = FileValidator(settings.max_file_mb)
        self.parser = DocumentParser(image_describer)
        self.chunker = ParentChildChunker(settings.child_tokens, settings.parent_tokens, settings.overlap_ratio)

    def ingest(self, path: str | Path, *, source_id: str | None = None, owner: str = "admin",
               acl: list[Role] | None = None) -> dict[str, Any]:
        source = Path(path).resolve()
        self.validator.validate(source)
        digest = self.validator.sha256(source)
        source_id = source_id or source.name
        acl = acl or [Role.GUEST, Role.OPERATOR, Role.ADMIN]
        safe_source = hashlib.sha256(source_id.encode()).hexdigest()[:16]
        stored_dir = self.settings.documents_dir / safe_source / digest
        stored_dir.mkdir(parents=True, exist_ok=True)
        stored_path = stored_dir / source.name
        if not stored_path.exists():
            shutil.copy2(source, stored_path)
        version, duplicate = self.database.begin_document_version(
            source_id=source_id, filename=source.name, source_type=source.suffix.lower().lstrip("."),
            sha256=digest, owner=owner, acl=acl, parser_version=self.PARSER_VERSION,
            stored_path=str(stored_path),
        )
        if duplicate:
            return {"document": version.model_dump(mode="json"), "duplicate": True}
        try:
            parsed, products = self.parser.parse(stored_path)
            section: list[str] = []
            elements: list[DocumentElement] = []
            for order, item in enumerate(parsed):
                if item.content_type == "title":
                    section = item.section_path or [item.title or item.text]
                element_id = hashlib.sha256(
                    f"{version.document_id}:{version.version}:{order}:{item.text}".encode()
                ).hexdigest()[:32]
                elements.append(DocumentElement(
                    element_id=element_id, document_id=version.document_id,
                    document_version=version.version, order=order, content_type=item.content_type,
                    text=item.text, title=item.title, section_path=item.section_path or section,
                    page_no=item.page_no, sheet_name=item.sheet_name, cell_range=item.cell_range,
                    bbox=item.bbox, table_json=item.table_json, image_ref=item.image_ref,
                    metadata={"mime": mimetypes.guess_type(source.name)[0], **item.metadata},
                ))
            for product in products:
                product.source_document_id = version.document_id
                product.source_document_version = version.version
                product.acl = acl
            chunks = self.chunker.chunk(version.document_id, version.version, acl, elements)
            self.database.complete_document_version(version, elements, chunks, products)
            return {
                "document": {**version.model_dump(mode="json"), "status": "ready", "active": True},
                "duplicate": False, "elements": len(elements), "chunks": len(chunks),
                "products": len(products), "warnings": self._warnings(parsed),
            }
        except Exception as exc:
            self.database.fail_document_version(version.document_id, version.version, str(exc))
            raise

    @staticmethod
    def _warnings(elements: list[ParsedElement]) -> list[str]:
        warnings = []
        if not elements:
            warnings.append("文档未解析出可索引内容")
        if any(any(pattern.search(item.text) for pattern in INJECTION_PATTERNS) for item in elements):
            warnings.append("检测到疑似提示注入内容，相关块默认不参与检索")
        return warnings
