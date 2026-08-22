"""
Data preprocessing pipeline for the unmanned vending vehicle RAG project.

It converts heterogeneous source files into a normalized JSONL corpus that can
be embedded and indexed later. Supported inputs:
- docx/pdf/txt/md for long documents such as school history and lab brochures
- xlsx/csv/json for product catalogs

Example:
    python rag_data_preprocessor.py --input ./data/raw --output ./dataset/preprocessed
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional


try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover
    DocxDocument = None

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover
    PdfReader = None

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


@dataclass
class SourceDocument:
    doc_id: str
    title: str
    source_path: str
    source_type: str
    knowledge_type: str
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass
class Chunk:
    chunk_id: str
    parent_id: str
    title: str
    content: str
    source_path: str
    source_type: str
    knowledge_type: str
    chunk_index: int
    metadata: Dict[str, Any] = field(default_factory=dict)


class RAGDataPreprocessor:
    def __init__(self, chunk_size: int = 700, chunk_overlap: int = 120):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def preprocess_dir(self, input_dir: Path, output_dir: Path) -> Dict[str, Any]:
        input_dir = input_dir.resolve()
        output_dir.mkdir(parents=True, exist_ok=True)

        docs: List[SourceDocument] = []
        for file_path in sorted(self._iter_supported_files(input_dir)):
            docs.extend(self.load_file(file_path, input_dir))

        chunks: List[Chunk] = []
        for doc in docs:
            chunks.extend(self.chunk_document(doc))

        self._write_jsonl(output_dir / "documents.jsonl", [asdict(doc) for doc in docs])
        self._write_jsonl(output_dir / "chunks.jsonl", [asdict(chunk) for chunk in chunks])

        manifest = self._build_manifest(input_dir, docs, chunks)
        (output_dir / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return manifest

    def load_file(self, file_path: Path, input_root: Path) -> List[SourceDocument]:
        suffix = file_path.suffix.lower()
        knowledge_type = self._infer_knowledge_type(file_path)

        if suffix == ".docx":
            content = self._read_docx(file_path)
            return [self._make_doc(file_path, input_root, knowledge_type, content)]
        if suffix == ".pdf":
            content = self._read_pdf(file_path)
            return [self._make_doc(file_path, input_root, knowledge_type, content)]
        if suffix in {".txt", ".md"}:
            content = file_path.read_text(encoding="utf-8", errors="ignore")
            return [self._make_doc(file_path, input_root, knowledge_type, content)]
        if suffix == ".xlsx":
            return self._read_xlsx_products(file_path, input_root)
        if suffix == ".csv":
            return self._read_csv_products(file_path, input_root)
        if suffix == ".json":
            return self._read_json_products(file_path, input_root)

        return []

    def chunk_document(self, doc: SourceDocument) -> List[Chunk]:
        if doc.knowledge_type == "product":
            return [
                Chunk(
                    chunk_id=self._stable_id(f"{doc.doc_id}:0:{doc.content}"),
                    parent_id=doc.doc_id,
                    title=doc.title,
                    content=doc.content,
                    source_path=doc.source_path,
                    source_type=doc.source_type,
                    knowledge_type=doc.knowledge_type,
                    chunk_index=0,
                    metadata=doc.metadata,
                )
            ]

        sections = self._split_by_headings(doc.content)
        chunks: List[Chunk] = []
        for section_title, section_text in sections:
            for piece in self._sliding_window(section_text):
                chunk_index = len(chunks)
                chunk_title = section_title or doc.title
                chunks.append(
                    Chunk(
                        chunk_id=self._stable_id(f"{doc.doc_id}:{chunk_index}:{piece}"),
                        parent_id=doc.doc_id,
                        title=chunk_title,
                        content=piece,
                        source_path=doc.source_path,
                        source_type=doc.source_type,
                        knowledge_type=doc.knowledge_type,
                        chunk_index=chunk_index,
                        metadata={**doc.metadata, "section_title": chunk_title},
                    )
                )
        return chunks

    def _make_doc(
        self, file_path: Path, input_root: Path, knowledge_type: str, raw_content: str
    ) -> SourceDocument:
        relative_path = file_path.resolve().relative_to(input_root.resolve()).as_posix()
        content = self._clean_text(raw_content)
        title = file_path.stem
        return SourceDocument(
            doc_id=self._stable_id(relative_path),
            title=title,
            source_path=relative_path,
            source_type=file_path.suffix.lower().lstrip("."),
            knowledge_type=knowledge_type,
            content=content,
            metadata=self._extract_metadata(title, content, knowledge_type),
        )

    def _read_docx(self, file_path: Path) -> str:
        if DocxDocument is None:
            raise RuntimeError("Please install python-docx to read .docx files.")
        doc = DocxDocument(str(file_path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        table_rows = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    table_rows.append(" | ".join(cells))
        return "\n".join(paragraphs + table_rows)

    def _read_pdf(self, file_path: Path) -> str:
        if PdfReader is None:
            raise RuntimeError("Please install pypdf to read .pdf files.")
        reader = PdfReader(str(file_path))
        pages = []
        for page_no, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"\n# Page {page_no}\n{text}")
        return "\n".join(pages)

    def _read_xlsx_products(self, file_path: Path, input_root: Path) -> List[SourceDocument]:
        if load_workbook is None:
            raise RuntimeError("Please install openpyxl to read .xlsx files.")
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        docs: List[SourceDocument] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [self._normalize_header(str(cell or "")) for cell in rows[0]]
            for row in rows[1:]:
                record = {
                    headers[i]: row[i]
                    for i in range(min(len(headers), len(row)))
                    if headers[i] and row[i] is not None
                }
                if record:
                    docs.append(self._product_doc(file_path, input_root, record, sheet.title))
        return docs

    def _read_csv_products(self, file_path: Path, input_root: Path) -> List[SourceDocument]:
        docs = []
        with file_path.open("r", encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                record = {self._normalize_header(k): v for k, v in row.items() if v}
                docs.append(self._product_doc(file_path, input_root, record))
        return docs

    def _read_json_products(self, file_path: Path, input_root: Path) -> List[SourceDocument]:
        data = json.loads(file_path.read_text(encoding="utf-8"))
        records = data if isinstance(data, list) else data.get("products", [])
        return [
            self._product_doc(file_path, input_root, {self._normalize_header(k): v for k, v in item.items()})
            for item in records
            if isinstance(item, dict)
        ]

    def _product_doc(
        self,
        file_path: Path,
        input_root: Path,
        record: Dict[str, Any],
        sheet_name: Optional[str] = None,
    ) -> SourceDocument:
        relative_path = file_path.resolve().relative_to(input_root.resolve()).as_posix()
        name = str(record.get("name") or record.get("product_name") or record.get("商品名称") or "未知商品")
        price = record.get("price") or record.get("商品价格") or record.get("价格")
        category = record.get("category") or record.get("商品类别") or record.get("类别") or "商品"
        effect = record.get("effect") or record.get("功效") or record.get("特点") or record.get("卖点") or ""
        recommendation = record.get("recommendation") or record.get("推荐语") or record.get("推荐话术") or ""
        crowd = record.get("target_user") or record.get("适合人群") or record.get("人群") or ""

        lines = [
            f"商品名称：{name}",
            f"类别：{category}",
            f"价格：{price}" if price not in (None, "") else "",
            f"功效/特点：{effect}" if effect else "",
            f"适合人群：{crowd}" if crowd else "",
            f"推荐语：{recommendation}" if recommendation else "",
        ]
        content = self._clean_text("\n".join(line for line in lines if line))
        metadata = {
            "product_name": name,
            "category": str(category),
            "price": price,
            "effect": str(effect),
            "recommendation": str(recommendation),
            "target_user": str(crowd),
            "sheet_name": sheet_name,
        }
        return SourceDocument(
            doc_id=self._stable_id(f"{relative_path}:{sheet_name or ''}:{name}:{price}"),
            title=name,
            source_path=relative_path,
            source_type=file_path.suffix.lower().lstrip("."),
            knowledge_type="product",
            content=content,
            metadata=metadata,
        )

    def _split_by_headings(self, text: str) -> List[tuple[str, str]]:
        lines = text.splitlines()
        sections: List[tuple[str, List[str]]] = []
        current_title = ""
        current_lines: List[str] = []

        for line in lines:
            stripped = line.strip()
            is_heading = (
                stripped.startswith("#")
                or re.match(r"^第[一二三四五六七八九十百0-9]+[章节篇编].{0,40}$", stripped)
                or re.match(r"^[0-9]+[\.、]\s*.{1,40}$", stripped)
            )
            if is_heading and current_lines:
                sections.append((current_title, current_lines))
                current_lines = []
                current_title = stripped.lstrip("#").strip()
            elif is_heading:
                current_title = stripped.lstrip("#").strip()
            else:
                current_lines.append(line)

        if current_lines:
            sections.append((current_title, current_lines))

        return [(title, self._clean_text("\n".join(body))) for title, body in sections if body]

    def _sliding_window(self, text: str) -> Iterable[str]:
        text = self._clean_text(text)
        if len(text) <= self.chunk_size:
            yield text
            return

        start = 0
        while start < len(text):
            end = min(start + self.chunk_size, len(text))
            window = text[start:end]
            sentence_end = max(window.rfind("。"), window.rfind("！"), window.rfind("？"), window.rfind("\n"))
            if sentence_end > self.chunk_size * 0.55:
                end = start + sentence_end + 1
            yield text[start:end].strip()
            if end >= len(text):
                break
            start = max(end - self.chunk_overlap, start + 1)

    def _extract_metadata(self, title: str, content: str, knowledge_type: str) -> Dict[str, Any]:
        metadata = {
            "keywords": self._keywords_from_text(f"{title}\n{content}"),
            "char_count": len(content),
        }
        if "北洋" in title or "校史" in title:
            metadata["domain"] = "school_history"
        elif "课题组" in title or "先进动力" in title or "车辆智能控制" in title:
            metadata["domain"] = "lab_profile"
        else:
            metadata["domain"] = knowledge_type
        return metadata

    def _infer_knowledge_type(self, file_path: Path) -> str:
        name = file_path.as_posix().lower()
        if any(key in name for key in ["product", "商品", "售卖", "price", "goods"]):
            return "product"
        if any(key in name for key in ["北洋", "校史", "history"]):
            return "school_history"
        if any(key in name for key in ["课题组", "宣传册", "lab", "先进动力", "车辆智能控制"]):
            return "lab_profile"
        return "general"

    def _keywords_from_text(self, text: str) -> List[str]:
        candidates = [
            "北洋大学", "天津大学", "校史", "先进动力", "车辆智能控制", "无人驾驶",
            "售卖车", "商品", "价格", "推荐", "饮品", "零食", "文创", "实验室",
        ]
        return [word for word in candidates if word in text]

    def _clean_text(self, text: str) -> str:
        text = text.replace("\u3000", " ").replace("\xa0", " ")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r"([。！？；])\s+", r"\1\n", text)
        return text.strip()

    def _normalize_header(self, header: str) -> str:
        header = header.strip()
        mapping = {
            "名称": "name",
            "商品名": "name",
            "商品名称": "name",
            "价格": "price",
            "单价": "price",
            "类别": "category",
            "分类": "category",
            "功效": "effect",
            "特点": "effect",
            "卖点": "effect",
            "推荐语": "recommendation",
            "推荐话术": "recommendation",
            "适合人群": "target_user",
            "人群": "target_user",
        }
        return mapping.get(header, header)

    def _iter_supported_files(self, input_dir: Path) -> Iterable[Path]:
        supported = {".docx", ".pdf", ".xlsx", ".csv", ".json", ".txt", ".md"}
        for path in input_dir.rglob("*"):
            if path.is_file() and path.suffix.lower() in supported and not path.name.startswith("~$"):
                yield path

    def _build_manifest(
        self, input_dir: Path, docs: List[SourceDocument], chunks: List[Chunk]
    ) -> Dict[str, Any]:
        by_type: Dict[str, int] = {}
        by_knowledge: Dict[str, int] = {}
        for doc in docs:
            by_type[doc.source_type] = by_type.get(doc.source_type, 0) + 1
            by_knowledge[doc.knowledge_type] = by_knowledge.get(doc.knowledge_type, 0) + 1

        return {
            "input_dir": str(input_dir),
            "document_count": len(docs),
            "chunk_count": len(chunks),
            "chunk_size": self.chunk_size,
            "chunk_overlap": self.chunk_overlap,
            "source_type_distribution": by_type,
            "knowledge_type_distribution": by_knowledge,
            "outputs": ["documents.jsonl", "chunks.jsonl", "manifest.json"],
        }

    def _write_jsonl(self, path: Path, rows: List[Dict[str, Any]]) -> None:
        with path.open("w", encoding="utf-8") as f:
            for row in rows:
                f.write(json.dumps(row, ensure_ascii=False) + "\n")

    def _stable_id(self, text: str) -> str:
        return hashlib.md5(text.encode("utf-8")).hexdigest()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Preprocess RAG source files into JSONL chunks.")
    parser.add_argument("--input", required=True, help="Raw data directory.")
    parser.add_argument("--output", default="./dataset/preprocessed", help="Output directory.")
    parser.add_argument("--chunk-size", type=int, default=700)
    parser.add_argument("--chunk-overlap", type=int, default=120)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    preprocessor = RAGDataPreprocessor(
        chunk_size=args.chunk_size,
        chunk_overlap=args.chunk_overlap,
    )
    manifest = preprocessor.preprocess_dir(Path(args.input), Path(args.output))
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
